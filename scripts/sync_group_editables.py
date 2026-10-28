"""Keep downloadable group worksheet editing formats aligned with the revised PDF."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import subprocess
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'scripts/templates'
DEST = ROOT / 'materials'
REPLACE = {
    '중요한 이유': '장단점',
    '02   의견 모으기': '02   의견 모으기 · 장단점 검토하기',
    '의견을 비교하며 나온 좋은 점과 걱정되는 점': '장점: 좋은 점과 그 이유',
    '우리 모둠의 최종 제안과 선택한 이유': '우리 모둠의 최종 의견과 근거',
}

doc = Document(BASE / 'group-11-12.docx')
for paragraph in [*doc.paragraphs, *(p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs)]:
    if paragraph.text in REPLACE:
        text = REPLACE[paragraph.text]
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
review, final, removed = doc.tables[2:5]
old_width = review.columns[0].width
review.columns[0].width = int(old_width / 2)
left = review.cell(0, 0)
left.width = int(old_width / 2)
review.add_column(int(old_width / 2))
right = review.cell(0, 1)
# Copy the paragraph and cell style instead of using the document's defaults.
right._tc.replace(right.paragraphs[0]._p, deepcopy(left.paragraphs[0]._p))
right.paragraphs[0].runs[0].text = '단점: 걱정되는 점과 그 이유'
left_props = left._tc.tcPr
right_props = right._tc.tcPr
for child in left_props:
    if child.tag.rsplit('}', 1)[-1] != 'tcW':
        right_props.append(deepcopy(child))
final.rows[0].height = Pt(101)
node = removed._tbl
previous = node.getprevious()
node.getparent().remove(node)
if previous is not None and previous.tag.endswith('}p') and not ''.join(previous.itertext()).strip():
    previous.getparent().remove(previous)

agreements = ['발언권 얻고 말하기', '의견을 끝까지 듣기', '이유와 근거를 들어 말하기', '서로 다른 의견을 존중하기', '충분히 협의하여 결정하기']
doc.add_page_break()
def back_paragraph(text, size, bold=False, after=8):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = '맑은 고딕'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    run.font.size = Pt(size)
    run.bold = bold
    return paragraph
back_paragraph('우리 지역 보행 안전 토의   /   11~12차시', 9, True)
back_paragraph('우리 모둠의 토의 약속', 22, True, 16)
for text in agreements:
    back_paragraph('• '+text, 13, after=10)
back_paragraph('자유 메모', 12.5, True, 7)
memo = doc.add_table(rows=1, cols=1)
memo.autofit = False
memo.columns[0].width = old_width
memo.cell(0, 0).width = old_width
memo.rows[0].height = Pt(425)
memo.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
props = memo.cell(0, 0)._tc.get_or_add_tcPr()
borders = OxmlElement('w:tcBorders')
for side in ['top', 'left', 'bottom', 'right']:
    edge = OxmlElement('w:'+side)
    for key, value in [('val','single'),('sz','5'),('color','BDC6CD')]:
        edge.set(qn('w:'+key), value)
    borders.append(edge)
props.append(borders)
doc.add_paragraph().paragraph_format.space_after = Pt(0)
doc.save(DEST / 'group-11-12.docx')

with ZipFile(BASE / 'group-11-12.hwpx') as z:
    files = [(info, z.read(info.filename)) for info in z.infolist()]
section = etree.fromstring(dict((i.filename, data) for i, data in files)['Contents/section0.xml'])
hp = {'hp': 'http://www.hancom.co.kr/hwpml/2011/paragraph'}
for text in section.findall('.//hp:t', hp):
    if text.text in REPLACE:
        text.text = REPLACE[text.text]
tables = section.findall('.//hp:tbl', hp)
review, final, removed = tables[2:5]
review.set('colCnt', '2')
row = review.find('hp:tr', hp)
left = row.find('hp:tc', hp)
right = deepcopy(left)
new_id = 1 + max(int(p.get('id', '0')) for p in section.findall('.//hp:p', hp))
for p in right.findall('.//hp:p', hp):
    p.set('id', str(new_id))
    new_id += 1
for cell, number, text in [(left, 0, '장점: 좋은 점과 그 이유'), (right, 1, '단점: 걱정되는 점과 그 이유')]:
    cell.find('hp:cellAddr', hp).set('colAddr', str(number))
    cell.find('hp:cellSz', hp).set('width', str(25227 if number == 0 else 25228))
    cell.find('.//hp:t', hp).text = text
row.append(right)
final.find('hp:sz', hp).set('height', '10100')
final.find('.//hp:cellSz', hp).set('height', '10100')
paragraph = removed.getparent().getparent()
assert paragraph.tag.endswith('}p')
paragraph.getparent().remove(paragraph)
# Build the back side from existing paragraph and table styles.
last_id = max(int(p.get('id', '0')) for p in section.findall('.//hp:p', hp))
def hwp_paragraph(text, model, page_break=False):
    global last_id
    paragraph = deepcopy(model)
    last_id += 1
    paragraph.set('id', str(last_id))
    paragraph.set('pageBreak', '1' if page_break else '0')
    for child in list(paragraph):
        if child.tag != '{'+hp['hp']+'}run':
            paragraph.remove(child)
    runs = paragraph.findall('hp:run', hp)
    for run in runs[1:]:
        paragraph.remove(run)
    run = runs[0]
    for child in list(run):
        run.remove(child)
    etree.SubElement(run, '{'+hp['hp']+'}t').text = text
    return paragraph
top_paragraphs = section.findall('hp:p', hp)
title_model = next(p for p in top_paragraphs if ''.join(p.itertext()) == '모둠 토의하기')
body_model = next(p for p in top_paragraphs if '모둠에서 한 장을' in ''.join(p.itertext()))
heading_model = next(p for p in top_paragraphs if '01   의견 나누기' in ''.join(p.itertext()))
section.append(hwp_paragraph('우리 모둠의 토의 약속', title_model, page_break=True))
for text in agreements:
    section.append(hwp_paragraph('• '+text, body_model))
section.append(hwp_paragraph('자유 메모', heading_model))
memo_paragraph = deepcopy(final.getparent().getparent())
last_id += 1
memo_paragraph.set('id', str(last_id))
memo_table = memo_paragraph.find('.//hp:tbl', hp)
memo_table.set('id', '6')
memo_table.find('hp:sz', hp).set('height', '42500')
memo_table.find('.//hp:cellSz', hp).set('height', '42500')
for p in memo_table.findall('.//hp:p', hp):
    last_id += 1
    p.set('id', str(last_id))
for t in memo_table.findall('.//hp:t', hp):
    t.text = None
section.append(memo_paragraph)
section_bytes = etree.tostring(section, xml_declaration=True, encoding='UTF-8', standalone=True)
preview = subprocess.run(['pdftoppm', '-scale-to', '700', '-singlefile', '-png', str(DEST / 'group-11-12.pdf')], capture_output=True, check=True).stdout
with ZipFile(DEST / 'group-11-12.hwpx', 'w') as out:
    for info, data in files:
        if info.filename == 'Contents/section0.xml':
            data = section_bytes
        elif info.filename == 'Preview/PrvText.txt':
            data = '\n'.join(t.text for t in section.findall('.//hp:t', hp) if t.text).encode('utf-8')
        elif info.filename == 'Preview/PrvImage.png':
            data = preview
        out.writestr(info, data)
with ZipFile(DEST / 'group-11-12.hwpx') as check:
    assert check.testzip() is None
    text = check.read('Contents/section0.xml').decode('utf-8')
    assert '친구 의견을 반영해 고치거나 더한 점' not in text
    assert '우리 모둠의 최종 의견과 근거' in text
print('Group DOCX and HWPX updated.')
